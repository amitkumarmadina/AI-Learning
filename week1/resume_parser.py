import os
import json
import time
from pathlib import Path

from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document


# ============================================================
# 1. ENVIRONMENT / GROQ CLIENT
# ============================================================

load_dotenv()

my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("GROQ_API_KEY not found")

client = Groq(api_key=my_api_key)

model = "llama-3.1-8b-instant"


# ============================================================
# 2. JOB DESCRIPTION
# ============================================================

job_description = f"""
Description
"This role is intended for 2026 and 2025 graduates only."

At Amazon, we hire the best minds in technology to innovate and build
on behalf of our customers. The focus we have on our customers is why
we are one of the world’s most beloved brands – customer obsession is
part of our company DNA.

Our Software Development Engineers (SDEs) use cutting-edge technology
to solve complex problems and get to see the impact of their work
first-hand.

The challenges SDEs solve at Amazon are big and influence millions of
customers, sellers, and products around the world.

We are looking for individuals who are passionate about creating new
products, features, and services from scratch while managing ambiguity
and the pace of a company where development cycles are measured in
weeks, not years.

Key job responsibilities

- Collaborate with experienced cross-disciplinary Amazonians to conceive,
  design, and bring innovative products and services to market.
- Design and build innovative technologies in a large distributed
  computing environment and help lead fundamental changes in the industry.
- Create solutions to run predictions on distributed systems with exposure
  to innovative technologies at incredible scale and speed.
- Build distributed storage, index, and query systems that are scalable,
  fault-tolerant, low cost, and easy to manage/use.
- Design and code the right solutions starting with broadly defined problems.
- Work in an agile environment to deliver high-quality software.

Basic Qualifications

- Bachelor's degree or above in computer science, computer engineering,
  or related field.
- Knowledge of Computer Science fundamentals such as object-oriented
  design, algorithm design, data structures, problem solving, and
  complexity analysis.
- Knowledge of programming languages such as C/C++, Python, Java or Perl.

Preferred Qualifications

- Previous technical internship(s).
- Experience with distributed, multi-tiered systems, algorithms,
  and relational databases.
- Experience in optimization mathematics such as linear programming
  and nonlinear optimization.
- Effectively articulate technical challenges and solutions.
- Adept at handling ambiguous or undefined problems as well as ability
  to think abstractly.
f"""


# ============================================================
# 3. PYDANTIC MODEL FOR JOB
# ============================================================

class JobD(BaseModel):
    role: str
    required_skills: list[str] = Field(default_factory=list)
    preferred_skills: list[str] = Field(default_factory=list)
    minimum_experience: float | None = None
    educational_requirements: list[str] = Field(default_factory=list)
    responsibilities: list[str] = Field(default_factory=list)


# ============================================================
# 4. EXTRACT JOB INFORMATION USING LLM
# ============================================================

jobd_schema = JobD.model_json_schema()

job_system_prompt = f"""
You are an expert HR assistant.

Extract structured information from the job description.

Return ONLY valid JSON.

The JSON MUST use EXACTLY these keys:

- role
- required_skills
- preferred_skills
- minimum_experience
- educational_requirements
- responsibilities

Do NOT capitalize the keys.
Do NOT rename the keys.
Do NOT use spaces in the keys.
Do NOT add extra keys.

Schema:
{jobd_schema}

Rules:
1. If minimum experience is not mentioned, return null.
2. If information for a list is missing, return [].
3. Do not invent information.
4. preferred_skills must be spelled exactly as "preferred_skills".
f"""

job_user_prompt = f"""
Analyse the following job description:

{job_description}
f"""

job_messages = [
    {
        "role": "system",
        "content": job_system_prompt
    },
    {
        "role": "user",
        "content": job_user_prompt
    }
]

response_format = {
    "type": "json_object"
}

response = client.chat.completions.create(
    model=model,
    messages=job_messages,
    response_format=response_format
)

raw_job_json = response.choices[0].message.content

if not raw_job_json:
    raise ValueError("LLM returned an empty job response")

job_data = json.loads(raw_job_json)

job = JobD(**job_data)

print("Job parsed successfully.")
print("Role:", job.role)
print("Minimum experience:", job.minimum_experience)
print("Education:", job.educational_requirements)


# ============================================================
# 5. PYDANTIC MODELS FOR RESUME
# ============================================================

class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = Field(default_factory=list)


class Education(BaseModel):
    school_name: str | None = None
    graduation: str | None = None


class Project(BaseModel):
    name: str | None = None
    start_date: str | None = None


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    total_experience_years: float | None = None

    skills: list[str] = Field(default_factory=list)

    experience: list[Experience] = Field(default_factory=list)

    education: list[Education] = Field(default_factory=list)

    projects: list[Project] = Field(default_factory=list)

    certifications: list[str] = Field(default_factory=list)


# ============================================================
# 6. MATCH RESULT MODEL
# ============================================================

class MatchResult(BaseModel):
    score: float
    details: dict


# ============================================================
# 7. PARSE RESUME
# ============================================================

resume_schema = Resume.model_json_schema()


def parse_resume(resume_text):

    system_prompt = f"""
You are an expert resume parser.

Extract information from the resume based on its meaning,
not only based on exact section headings.

Different resumes may use different headings.

For example:

- Experience
- Professional Experience
- Work History
- Employment
- Internships

These may all contain relevant experience.

Skills may appear in:
- Skills section
- Work experience
- Internships
- Projects

Return ONLY valid JSON matching this schema:

{resume_schema}

IMPORTANT RULES:

1. Do not invent information.

2. If a single value is not available, return null.

3. If a list has no information, return [].

4. Include internships inside the experience list.

5. Extract skills mentioned across the entire resume.

6. Education MUST be an array of objects with:
   - school_name
   - graduation

7. Projects MUST be an array of objects with:
   - name
   - start_date

8. Experience MUST be an array of objects with:
   - company
   - role
   - duration
   - description
   - skills_used

9. total_experience_years must be a number.
   Example: 2, 3.5, 5

10. Do NOT return the JSON schema itself.

11. Do NOT return "properties", "type", or "required" as the top-level
    response.

12. Return the actual extracted resume data.
f"""

    user_prompt = f"""
Parse the following resume:

{resume_text}
f"""

    messages = [
        {
            "role": "system",
            "content": system_prompt
        },
        {
            "role": "user",
            "content": user_prompt
        }
    ]

    response_format = {
        "type": "json_object"
    }

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format=response_format
    )

    raw_output = response.choices[0].message.content

    if not raw_output:
        raise ValueError("LLM returned an empty resume response")

    data = json.loads(raw_output)

    print("\n========== RESUME DATA ==========")
    print(data)
    print("=================================")

    resume = Resume(**data)

    return resume


# ============================================================
# 8. READ PDF
# ============================================================

def read_pdf(file_path):

    reader = PdfReader(file_path)

    text = ""

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


# ============================================================
# 9. READ DOCX
# ============================================================

def read_docx(file_path):

    document = Document(file_path)

    text = ""

    # Normal paragraphs
    for paragraph in document.paragraphs:

        if paragraph.text.strip():
            text += paragraph.text + "\n"

    # Tables
    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                if cell.text.strip():
                    text += cell.text + "\n"

    return text


# ============================================================
# 10. READ RESUME BASED ON FILE TYPE
# ============================================================

def read_resume(file_path):

    extension = file_path.suffix.lower()

    if extension == ".pdf":

        return read_pdf(file_path)

    elif extension == ".docx":

        return read_docx(file_path)

    else:

        return None


# ============================================================
# 11. FINAL SCORE
# ============================================================

def final_score(job, resume):

    prompt = f"""
You are an expert HR recruiter.

Compare the candidate's resume with the job description.

JOB DESCRIPTION:
{job.model_dump_json(indent=2)}

CANDIDATE RESUME:
{resume.model_dump_json(indent=2)}

Return ONLY the actual result JSON.

IMPORTANT:
- Do NOT return the JSON schema.
- Do NOT return "properties".
- Do NOT return "type".
- Do NOT return "required".
- Do NOT return the schema itself.
- Return the actual comparison result.

Return EXACTLY this structure:

{{
    "score": 85,
    "details": {{
        "candidate_name": "Candidate name",
        "matching_skills": [
            "Python",
            "Data Structures"
        ],
        "missing_skills": [
            "Java"
        ],
        "meets_requirements": true,
        "final_verdict": "Strong candidate"
    }}
}}

Rules:

1. score must be a number from 0 to 100.

2. score represents the overall match percentage.

3. details must be an object.

4. candidate_name should come from the resume.

5. matching_skills should contain skills from the resume
   that are relevant to the job.

6. missing_skills should contain important job skills that
   are missing from the resume.

7. meets_requirements should be true or false.

8. final_verdict should be short.

9. Do not invent information.

10. Base the score on actual evidence from the resume and job description.
f"""

    messages = [
        {
            "role": "system",
            "content": "You are an expert HR recruiter."
        },
        {
            "role": "user",
            "content": prompt
        }
    ]

    response_format = {
        "type": "json_object"
    }

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format=response_format
    )

    raw_output = response.choices[0].message.content

    if not raw_output:
        raise ValueError("LLM returned an empty scoring response")

    data = json.loads(raw_output)

    print("\n========== MATCH RESULT ==========")
    print(data)
    print("==================================")

    result = MatchResult(**data)

    return result


# ============================================================
# 12. RESUME FOLDER
# ============================================================

resume_folder = Path("resumes")

if not resume_folder.exists():

    raise FileNotFoundError(
        f"Resume folder not found: {resume_folder.absolute()}\n"
        f"Create a folder named 'resumes' inside your project folder."
    )


# ============================================================
# 13. PROCESS ALL RESUMES
# ============================================================

all_results = []


for file_path in resume_folder.iterdir():

    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue

    print("\n======================================")
    print("Processing:", file_path.name)
    print("======================================")

    resume_text = read_resume(file_path)

    if not resume_text or not resume_text.strip():

        print("Could not extract text from:", file_path.name)
        continue

    # Parse resume
    parsed_resume = parse_resume(resume_text)

    # Small delay between API calls
    time.sleep(5)

    # Calculate score
    result = final_score(job, parsed_resume)

    time.sleep(5)

    print("\nScore:", result.score)

    print("Details:", result.details)

    all_results.append(
        {
            "name": parsed_resume.name or file_path.stem,
            "score": result.score,
            "details": result.details
        }
    )


# ============================================================
# 14. CHECK IF WE HAVE RESULTS
# ============================================================

if not all_results:

    print("\nNo valid resumes were processed.")

    raise SystemExit


# ============================================================
# 15. SORT CANDIDATES
# ============================================================

all_results.sort(
    key=lambda candidate: candidate["score"],
    reverse=True
)


# ============================================================
# 16. TOP 2 AND LOWEST 2
# ============================================================

top_2 = all_results[:2]

worst_2 = all_results[-2:]


# ============================================================
# 17. PRINT TOP 2
# ============================================================

print("\n\n======================================")
print("TOP 2 CANDIDATES")
print("======================================")


for candidate in top_2:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )

    print(candidate["details"])

    print("--------------------------------------")


# ============================================================
# 18. PRINT LOWEST 2
# ============================================================

print("\n\n======================================")
print("LOWEST 2 CANDIDATES")
print("======================================")


for candidate in worst_2:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )

    print(candidate["details"])

    print("--------------------------------------")